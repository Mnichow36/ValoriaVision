from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.services.counterparty_parser import extract_counterparty
from app.imports.csv_parser import RobustCsvParser, normalize_date, parse_amount, payload_to_json
from app.imports.import_models import ImportedTransaction, ImportAnalysis, ImportDiagnostic, ImportRejectedLine


class BankImportService:
    """Moteur d'import bancaire R7.2.

    Architecture : détection -> parsing -> normalisation -> diagnostic.
    Le CSV est traité par un moteur robuste/profilable. Les formats historiques
    XLSX/QIF/OFX restent supportés en non-régression et sont convertis vers le
    même modèle normalisé.
    """

    DATE_KEYS = ["date", "operation date", "date operation", "date opération", "date comptable"]
    LABEL_KEYS = ["libelle", "libellé", "label", "description", "operation", "opération", "memo", "name"]
    AMOUNT_KEYS = ["amount", "montant", "montant eur", "valeur"]
    DEBIT_KEYS = ["debit", "débit", "sortie"]
    CREDIT_KEYS = ["credit", "crédit", "entree", "entrée"]

    def __init__(self) -> None:
        self.csv_parser = RobustCsvParser()
        self.last_analysis: ImportAnalysis | None = None

    def analyze(self, path: str | Path) -> ImportAnalysis:
        file_path = Path(path)
        suffix = file_path.suffix.lower()
        if suffix == ".csv":
            analysis = self.csv_parser.parse(file_path)
        elif suffix == ".xlsx":
            rows = self.parse_xlsx(file_path)
            analysis = ImportAnalysis(file_name=file_path.name, valid_transactions=rows)
            analysis.diagnostic = ImportDiagnostic(stage="normalisation", format_name="XLSX", profile_name="XLSX générique", message=f"{len(rows)} opération(s) reconnue(s).")
        elif suffix == ".qif":
            rows = self.parse_qif(file_path)
            analysis = ImportAnalysis(file_name=file_path.name, valid_transactions=rows)
            analysis.diagnostic = ImportDiagnostic(stage="normalisation", format_name="QIF", profile_name="QIF générique", message=f"{len(rows)} opération(s) reconnue(s).")
        elif suffix == ".ofx":
            rows = self.parse_ofx(file_path)
            analysis = ImportAnalysis(file_name=file_path.name, valid_transactions=rows)
            analysis.diagnostic = ImportDiagnostic(stage="normalisation", format_name="OFX", profile_name="OFX générique", message=f"{len(rows)} opération(s) reconnue(s).")
        else:
            analysis = ImportAnalysis(file_name=file_path.name)
            analysis.diagnostic = ImportDiagnostic(stage="format", format_name=suffix, message=f"Format non supporté : {suffix}")
        self.identify_counterparties(analysis)
        self.last_analysis = analysis
        return analysis

    @staticmethod
    def identify_counterparties(analysis: ImportAnalysis) -> ImportAnalysis:
        for transaction in analysis.valid_transactions:
            transaction.beneficiary = extract_counterparty(transaction.label)
        return analysis

    def parse(self, path: str | Path) -> list[ImportedTransaction]:
        return self.analyze(path).valid_transactions

    def parse_xlsx(self, path: Path) -> list[ImportedTransaction]:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(cell or "").strip() for cell in rows[0]]
        dict_rows: list[dict[str, Any]] = []
        for raw in rows[1:]:
            dict_rows.append({headers[i]: raw[i] if i < len(raw) else None for i in range(len(headers))})
        return self.rows_to_transactions(dict_rows, profile="XLSX générique")

    def parse_qif(self, path: Path) -> list[ImportedTransaction]:
        txs: list[ImportedTransaction] = []
        current: dict[str, str] = {}
        for line_no, line in enumerate(path.read_text(encoding="utf-8-sig", errors="replace").splitlines(), start=1):
            line = line.strip()
            if not line:
                continue
            if line == "^":
                if current:
                    tx_date = normalize_date(current.get("D", ""))
                    amount = parse_amount(current.get("T", "0"))
                    if tx_date and amount is not None:
                        txs.append(
                            ImportedTransaction(
                                tx_date=tx_date,
                                label=current.get("P") or current.get("M") or "Opération QIF",
                                amount=amount,
                                source_line=line_no,
                                source_profile="QIF générique",
                                raw_payload=dict(current),
                            )
                        )
                current = {}
            elif len(line) >= 2:
                current[line[0]] = line[1:]
        return txs

    def parse_ofx(self, path: Path) -> list[ImportedTransaction]:
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        blocks = re.findall(r"<STMTTRN>(.*?)(?=<STMTTRN>|</BANKTRANLIST>)", text, flags=re.I | re.S)
        txs: list[ImportedTransaction] = []
        for index, block in enumerate(blocks, start=1):
            date_raw = self.tag_value(block, "DTPOSTED")
            amount_raw = self.tag_value(block, "TRNAMT")
            name = self.tag_value(block, "NAME") or self.tag_value(block, "MEMO") or "Opération OFX"
            tx_date = normalize_date(date_raw[:8])
            amount = parse_amount(amount_raw)
            if tx_date and amount is not None:
                txs.append(
                    ImportedTransaction(
                        tx_date=tx_date,
                        label=name.strip(),
                        amount=amount,
                        source_line=index,
                        source_profile="OFX générique",
                        raw_payload={"block_index": index, "name": name, "amount": amount_raw, "date": date_raw},
                    )
                )
        return txs

    @staticmethod
    def tag_value(block: str, tag: str) -> str:
        match = re.search(rf"<{tag}>([^<\r\n]+)", block, flags=re.I)
        return match.group(1).strip() if match else ""

    def rows_to_transactions(self, rows: list[dict[str, Any]], profile: str = "Générique") -> list[ImportedTransaction]:
        if not rows:
            return []
        keys = list(rows[0].keys())
        mapping = self.detect_mapping(keys)
        txs: list[ImportedTransaction] = []
        for line_no, row in enumerate(rows, start=2):
            try:
                date_raw = row.get(mapping.get("date", ""), "")
                label_raw = row.get(mapping.get("label", ""), "")
                tx_date = normalize_date(date_raw)
                label = str(label_raw or "").strip() or "Opération bancaire"
                amount = self.extract_amount(row, mapping)
                if tx_date and amount is not None:
                    txs.append(
                        ImportedTransaction(
                            tx_date=tx_date,
                            label=label,
                            amount=amount,
                            source_line=line_no,
                            source_profile=profile,
                            raw_payload={str(k): str(v) for k, v in row.items()},
                        )
                    )
            except Exception:
                continue
        return txs

    def detect_mapping(self, keys: list[str]) -> dict[str, str]:
        lowered = {self.clean_key(key): key for key in keys}
        return {
            "date": self.find_key(lowered, self.DATE_KEYS),
            "label": self.find_key(lowered, self.LABEL_KEYS),
            "amount": self.find_key(lowered, self.AMOUNT_KEYS),
            "debit": self.find_key(lowered, self.DEBIT_KEYS),
            "credit": self.find_key(lowered, self.CREDIT_KEYS),
        }

    @staticmethod
    def clean_key(key: str) -> str:
        return str(key or "").strip().lower().replace("_", " ")

    def find_key(self, lowered: dict[str, str], candidates: list[str]) -> str:
        for candidate in candidates:
            for clean, original in lowered.items():
                if candidate in clean:
                    return original
        return ""

    def extract_amount(self, row: dict[str, Any], mapping: dict[str, str]) -> float | None:
        if mapping.get("amount"):
            return parse_amount(row.get(mapping["amount"]))
        debit = parse_amount(row.get(mapping.get("debit", ""))) if mapping.get("debit") else 0.0
        credit = parse_amount(row.get(mapping.get("credit", ""))) if mapping.get("credit") else 0.0
        if debit and credit:
            return abs(credit) - abs(debit)
        if credit:
            return abs(credit)
        if debit:
            return -abs(debit)
        return None

    @staticmethod
    def to_amount(value: Any) -> float:
        parsed = parse_amount(value)
        return float(parsed or 0.0)

    @staticmethod
    def normalize_date(value: Any) -> str:
        return normalize_date(value)
